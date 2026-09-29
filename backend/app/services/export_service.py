import csv
import io
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.auth import User
from app.models.catalog import Product
from app.models.customers import Customer
from app.models.inventory import Inventory
from app.models.purchasing import PurchaseOrder, Supplier
from app.models.sales import Sale
from app.services.report_service import Report, col

MAX_EXPORT_ROWS = 50_000


def _cell(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    return v


def _csv_safe(v: Any) -> Any:
    """Stop spreadsheet formula injection from user-controlled text (=, +, -, @ prefixes)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def to_csv(report: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c["label"] for c in report.columns])
    for row in report.rows:
        w.writerow([_csv_safe(_cell(row.get(c["key"], ""))) for c in report.columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads Bangla text correctly


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31]
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    if report.period:
        ws.append([f"Period: {report.period['start']} to {report.period['end']}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c["label"] for c in report.columns])
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(horizontal="center")
    for row in report.rows:
        ws.append([_csv_safe(_cell(row.get(c["key"], ""))) for c in report.columns])
    for idx, c in enumerate(report.columns, start=1):
        letter = get_column_letter(idx)
        if c["type"] in ("money", "qty", "percent"):
            for cell in ws[letter][header_row:]:
                cell.number_format = "#,##0.00"
        width = max([len(str(c["label"]))] + [len(str(r.get(c["key"], ""))) for r in report.rows[:200]]) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), 48)
    if report.summary:
        ws.append([])
        for k, v in report.summary.items():
            ws.append([k.replace("_", " ").title(), _cell(v)])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


MEDIA = {"csv": "text/csv; charset=utf-8", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


def render(report: Report, fmt: str) -> tuple[bytes, str]:
    return (to_csv(report) if fmt == "csv" else to_xlsx(report)), MEDIA[fmt]


# ---- entity datasets ------------------------------------------------------------------

def dataset(db: Session, name: str, user: User) -> Report:
    codes = user.permission_codes
    if name == "products":
        cost = "product.cost" in codes
        cols = [col("sku", "SKU"), col("barcode", "Barcode"), col("name", "Name"), col("name_bn", "Bangla name"), col("category", "Category"),
                col("brand", "Brand"), col("unit", "Unit")]
        if cost:
            cols.append(col("purchase_price", "Purchase price", "money"))
        cols += [col("selling_price", "Selling price", "money"), col("mrp", "MRP", "money"), col("reorder_level", "Reorder level", "qty"),
                 col("min_stock", "Min stock", "qty"), col("stock", "In stock", "qty"), col("track_expiry", "Track expiry"), col("active", "Active")]
        rows = []
        for p in db.scalars(select(Product).where(Product.is_deleted.is_(False)).order_by(Product.name).limit(MAX_EXPORT_ROWS)).unique():
            rows.append({"sku": p.sku, "barcode": p.barcode or "", "name": p.name, "name_bn": p.name_bn or "", "category": p.category.name if p.category else "",
                         "brand": p.brand.name if p.brand else "", "unit": p.unit.name, "purchase_price": p.purchase_price, "selling_price": p.selling_price,
                         "mrp": p.mrp or "", "reorder_level": p.reorder_level, "min_stock": p.min_stock, "stock": p.current_stock,
                         "track_expiry": "yes" if p.track_expiry else "no", "active": "yes" if p.is_active else "no"})
        return Report("Products", cols, rows)
    if name == "inventory":
        from app.services import report_service
        from app.utils.dates import range_bounds
        from datetime import date

        f = report_service.Filters(lo=None, hi=None, start=date.today(), end=date.today(), tz=None, limit=MAX_EXPORT_ROWS)  # type: ignore[arg-type]
        return report_service.current_inventory(db, f, "product.cost" in codes)
    if name == "customers":
        rows = [{"code": c.code, "name": c.name, "phone": c.phone or "", "email": c.email or "", "address": c.address or "", "type": c.customer_type,
                 "credit_limit": c.credit_limit, "balance": c.balance, "points": c.loyalty_points}
                for c in db.scalars(select(Customer).where(Customer.is_deleted.is_(False)).order_by(Customer.name).limit(MAX_EXPORT_ROWS))]
        return Report("Customers", [col("code", "Code"), col("name", "Name"), col("phone", "Phone"), col("email", "Email"), col("address", "Address"),
                                    col("type", "Type"), col("credit_limit", "Credit limit", "money"), col("balance", "Balance due", "money"),
                                    col("points", "Loyalty points", "int")], rows)
    if name == "suppliers":
        rows = [{"code": s.code, "name": s.name, "company": s.company or "", "phone": s.phone or "", "email": s.email or "", "contact": s.contact_person or "",
                 "terms": s.payment_terms_days, "balance": s.balance}
                for s in db.scalars(select(Supplier).where(Supplier.is_deleted.is_(False)).order_by(Supplier.name).limit(MAX_EXPORT_ROWS))]
        return Report("Suppliers", [col("code", "Code"), col("name", "Name"), col("company", "Company"), col("phone", "Phone"), col("email", "Email"),
                                    col("contact", "Contact person"), col("terms", "Terms (days)", "int"), col("balance", "Payable", "money")], rows)
    if name == "sales":
        stmt = select(Sale).order_by(Sale.id.desc()).limit(MAX_EXPORT_ROWS)
        if "sale.read_all" not in codes:
            stmt = stmt.where(Sale.cashier_id == user.id)
        rows = [{"invoice": s.invoice_number, "date": s.sale_date.strftime("%Y-%m-%d %H:%M"), "customer": s.customer.name if s.customer else "Walk-in",
                 "cashier": s.cashier.full_name, "subtotal": s.subtotal, "discount": s.discount_amount, "tax": s.tax_amount, "total": s.total_amount,
                 "paid": s.paid_amount, "due": s.due_amount, "status": s.status}
                for s in db.scalars(stmt).unique()]
        return Report("Sales", [col("invoice", "Invoice"), col("date", "Date (UTC)"), col("customer", "Customer"), col("cashier", "Cashier"),
                                col("subtotal", "Subtotal", "money"), col("discount", "Discount", "money"), col("tax", "VAT", "money"),
                                col("total", "Total", "money"), col("paid", "Paid", "money"), col("due", "Due", "money"), col("status", "Status")], rows)
    if name == "purchases":
        rows = [{"po": p.po_number, "date": p.order_date.isoformat(), "supplier": p.supplier.name, "status": p.status, "total": p.total_amount,
                 "received": p.received_value, "paid": p.paid_amount}
                for p in db.scalars(select(PurchaseOrder).order_by(PurchaseOrder.id.desc()).limit(MAX_EXPORT_ROWS)).unique()]
        return Report("Purchases", [col("po", "PO #"), col("date", "Date", "date"), col("supplier", "Supplier"), col("status", "Status"),
                                    col("total", "Total", "money"), col("received", "Received value", "money"), col("paid", "Paid", "money")], rows)
    from app.core.errors import NotFoundError

    raise NotFoundError(f"Unknown dataset '{name}'")


EXPORT_PERMISSIONS = {"products": "product.read", "inventory": "inventory.read", "customers": "customer.read", "suppliers": "supplier.read",
                      "sales": "sale.read", "purchases": "purchase.read"}
